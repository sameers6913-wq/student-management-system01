import customtkinter as ctk
from tkinter import messagebox, filedialog
from database import create_connection
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak
)

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from datetime import datetime


class ReportsPage(ctk.CTkFrame):

    def __init__(self, parent):
        super().__init__(parent, fg_color="#f5f7fb")

        self.create_header()
        self.create_summary_cards()
        self.create_export_section()
        self.create_charts()

        self.load_statistics()

    # =========================================================
    # HEADER
    # =========================================================

    def create_header(self):

        header = ctk.CTkFrame(
            self,
            fg_color="white",
            corner_radius=12
        )
        header.pack(
            fill="x",
            padx=20,
            pady=(20, 10)
        )

        title = ctk.CTkLabel(
            header,
            text="📊 Reports & Analytics",
            font=("Segoe UI", 26, "bold"),
            text_color="#1f2937"
        )
        title.pack(
            side="left",
            padx=20,
            pady=18
        )

        refresh_button = ctk.CTkButton(
            header,
            text="⟳ Refresh Reports",
            width=160,
            height=40,
            corner_radius=8,
            command=self.refresh_reports
        )
        refresh_button.pack(
            side="right",
            padx=20
        )

    # =========================================================
    # SUMMARY CARDS
    # =========================================================

    def create_summary_cards(self):

        cards_frame = ctk.CTkFrame(
            self,
            fg_color="transparent"
        )
        cards_frame.pack(
            fill="x",
            padx=20,
            pady=10
        )

        self.student_card = self.create_card(
            cards_frame,
            "👨‍🎓",
            "Total Students",
            "0"
        )

        self.teacher_card = self.create_card(
            cards_frame,
            "👨‍🏫",
            "Total Teachers",
            "0"
        )

        self.course_card = self.create_card(
            cards_frame,
            "📚",
            "Total Courses",
            "0"
        )

        self.subject_card = self.create_card(
            cards_frame,
            "📖",
            "Total Subjects",
            "0"
        )

        self.attendance_card = self.create_card(
            cards_frame,
            "📅",
            "Attendance",
            "0%"
        )

    def create_card(self, parent, icon, title, value):

        card = ctk.CTkFrame(
            parent,
            fg_color="white",
            corner_radius=12,
            height=110
        )

        card.pack(
            side="left",
            fill="both",
            expand=True,
            padx=6
        )

        icon_label = ctk.CTkLabel(
            card,
            text=icon,
            font=("Segoe UI", 24)
        )
        icon_label.pack(
            pady=(10, 0)
        )

        title_label = ctk.CTkLabel(
            card,
            text=title,
            font=("Segoe UI", 13),
            text_color="#6b7280"
        )
        title_label.pack()

        value_label = ctk.CTkLabel(
            card,
            text=value,
            font=("Segoe UI", 22, "bold"),
            text_color="#2563eb"
        )
        value_label.pack()

        return value_label

    # =========================================================
    # EXPORT SECTION
    # =========================================================

    def create_export_section(self):

        export_frame = ctk.CTkFrame(
            self,
            fg_color="white",
            corner_radius=12
        )
        export_frame.pack(
            fill="x",
            padx=20,
            pady=10
        )

        title = ctk.CTkLabel(
            export_frame,
            text="📥 Download Reports",
            font=("Segoe UI", 18, "bold"),
            text_color="#1f2937"
        )
        title.pack(
            side="left",
            padx=20,
            pady=15
        )

        pdf_button = ctk.CTkButton(
            export_frame,
            text="📄 Complete PDF",
            width=150,
            height=38,
            corner_radius=8,
            command=self.export_pdf
        )
        pdf_button.pack(
            side="right",
            padx=5
        )

        excel_button = ctk.CTkButton(
            export_frame,
            text="📊 Complete Excel",
            width=160,
            height=38,
            corner_radius=8,
            command=self.export_excel
        )
        excel_button.pack(
            side="right",
            padx=5
        )

        attendance_button = ctk.CTkButton(
            export_frame,
            text="📅 Attendance Excel",
            width=170,
            height=38,
            corner_radius=8,
            command=self.export_attendance_excel
        )
        attendance_button.pack(
            side="right",
            padx=5
        )

        marks_button = ctk.CTkButton(
            export_frame,
            text="📝 Marks Excel",
            width=140,
            height=38,
            corner_radius=8,
            command=self.export_marks_excel
        )
        marks_button.pack(
            side="right",
            padx=5
        )

    # =========================================================
    # CHART AREA
    # =========================================================

    def create_charts(self):

        charts_frame = ctk.CTkFrame(
            self,
            fg_color="transparent"
        )

        charts_frame.pack(
            fill="both",
            expand=True,
            padx=20,
            pady=10
        )

        # COURSE CHART

        course_frame = ctk.CTkFrame(
            charts_frame,
            fg_color="white",
            corner_radius=12
        )

        course_frame.pack(
            side="left",
            fill="both",
            expand=True,
            padx=(0, 8)
        )

        course_title = ctk.CTkLabel(
            course_frame,
            text="Students by Course",
            font=("Segoe UI", 18, "bold"),
            text_color="#1f2937"
        )
        course_title.pack(
            pady=10
        )

        self.course_chart_area = ctk.CTkFrame(
            course_frame,
            fg_color="white"
        )
        self.course_chart_area.pack(
            fill="both",
            expand=True,
            padx=10,
            pady=10
        )

        # GRADE CHART

        grade_frame = ctk.CTkFrame(
            charts_frame,
            fg_color="white",
            corner_radius=12
        )

        grade_frame.pack(
            side="right",
            fill="both",
            expand=True,
            padx=(8, 0)
        )

        grade_title = ctk.CTkLabel(
            grade_frame,
            text="Grade Distribution",
            font=("Segoe UI", 18, "bold"),
            text_color="#1f2937"
        )
        grade_title.pack(
            pady=10
        )

        self.grade_chart_area = ctk.CTkFrame(
            grade_frame,
            fg_color="white"
        )
        self.grade_chart_area.pack(
            fill="both",
            expand=True,
            padx=10,
            pady=10
        )

    # =========================================================
    # LOAD STATISTICS
    # =========================================================

    def load_statistics(self):

        connection = create_connection()

        if connection is None:
            messagebox.showerror(
                "Database Error",
                "Unable to connect to MySQL."
            )
            return

        try:

            cursor = connection.cursor()

            cursor.execute(
                "SELECT COUNT(*) FROM students"
            )
            total_students = cursor.fetchone()[0]

            cursor.execute(
                "SELECT COUNT(*) FROM teachers"
            )
            total_teachers = cursor.fetchone()[0]

            cursor.execute(
                "SELECT COUNT(*) FROM courses"
            )
            total_courses = cursor.fetchone()[0]

            cursor.execute(
                "SELECT COUNT(*) FROM subjects"
            )
            total_subjects = cursor.fetchone()[0]

            cursor.execute(
                """
                SELECT
                    SUM(
                        CASE
                            WHEN status = 'Present'
                            THEN 1
                            ELSE 0
                        END
                    ),
                    COUNT(*)
                FROM attendance
                """
            )

            attendance_result = cursor.fetchone()

            present_count = attendance_result[0] or 0
            total_attendance = attendance_result[1] or 0

            if total_attendance > 0:
                attendance_percentage = (
                    present_count / total_attendance
                ) * 100
            else:
                attendance_percentage = 0

            self.student_card.configure(
                text=str(total_students)
            )

            self.teacher_card.configure(
                text=str(total_teachers)
            )

            self.course_card.configure(
                text=str(total_courses)
            )

            self.subject_card.configure(
                text=str(total_subjects)
            )

            self.attendance_card.configure(
                text=f"{attendance_percentage:.1f}%"
            )

            self.load_course_chart(cursor)
            self.load_grade_chart(cursor)

            cursor.close()

        except Exception as e:

            messagebox.showerror(
                "Report Error",
                str(e)
            )

        finally:

            connection.close()

    # =========================================================
    # COURSE CHART
    # =========================================================

    def load_course_chart(self, cursor):

        for widget in self.course_chart_area.winfo_children():
            widget.destroy()

        cursor.execute(
            """
            SELECT
                c.course_name,
                COUNT(e.id)
            FROM courses c
            LEFT JOIN enrollments e
                ON c.course_code = e.course_code
            GROUP BY
                c.course_code,
                c.course_name
            ORDER BY COUNT(e.id) DESC
            """
        )

        data = cursor.fetchall()

        if not data:

            label = ctk.CTkLabel(
                self.course_chart_area,
                text="No enrollment data available",
                font=("Segoe UI", 15)
            )
            label.pack(expand=True)
            return

        course_names = [
            row[0] for row in data
        ]

        student_counts = [
            row[1] for row in data
        ]

        figure, axis = plt.subplots(
            figsize=(5, 4),
            dpi=90
        )

        axis.bar(
            course_names,
            student_counts
        )

        axis.set_title(
            "Students Enrolled by Course"
        )

        axis.set_ylabel(
            "Students"
        )

        axis.tick_params(
            axis="x",
            rotation=30
        )

        figure.tight_layout()

        canvas = FigureCanvasTkAgg(
            figure,
            master=self.course_chart_area
        )

        canvas.draw()

        canvas.get_tk_widget().pack(
            fill="both",
            expand=True
        )

        plt.close(figure)

    # =========================================================
    # GRADE CHART
    # =========================================================

    def load_grade_chart(self, cursor):

        for widget in self.grade_chart_area.winfo_children():
            widget.destroy()

        cursor.execute(
            """
            SELECT
                grade,
                COUNT(*)
            FROM marks
            WHERE grade IS NOT NULL
            AND grade != ''
            GROUP BY grade
            ORDER BY grade
            """
        )

        data = cursor.fetchall()

        if not data:

            label = ctk.CTkLabel(
                self.grade_chart_area,
                text="No marks data available",
                font=("Segoe UI", 15)
            )
            label.pack(expand=True)
            return

        grades = [
            row[0] for row in data
        ]

        counts = [
            row[1] for row in data
        ]

        figure, axis = plt.subplots(
            figsize=(5, 4),
            dpi=90
        )

        axis.pie(
            counts,
            labels=grades,
            autopct="%1.1f%%"
        )

        axis.set_title(
            "Grade Distribution"
        )

        figure.tight_layout()

        canvas = FigureCanvasTkAgg(
            figure,
            master=self.grade_chart_area
        )

        canvas.draw()

        canvas.get_tk_widget().pack(
            fill="both",
            expand=True
        )

        plt.close(figure)

    # =========================================================
    # DATABASE DATA
    # =========================================================

    def get_table_data(self, query):

        connection = create_connection()

        if connection is None:
            return None, None

        try:

            cursor = connection.cursor()
            cursor.execute(query)

            rows = cursor.fetchall()

            columns = [
                description[0]
                for description in cursor.description
            ]

            cursor.close()
            connection.close()

            return columns, rows

        except Exception as e:

            connection.close()

            messagebox.showerror(
                "Database Error",
                str(e)
            )

            return None, None

    # =========================================================
    # EXCEL SHEET HELPER
    # =========================================================

    def add_sheet_to_workbook(
        self,
        workbook,
        sheet_name,
        columns,
        rows
    ):

        sheet = workbook.create_sheet(
            sheet_name
        )

        for col_num, column_name in enumerate(
            columns,
            start=1
        ):

            cell = sheet.cell(
                row=1,
                column=col_num,
                value=column_name
            )

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        for row_num, row in enumerate(
            rows,
            start=2
        ):

            for col_num, value in enumerate(
                row,
                start=1
            ):

                sheet.cell(
                    row=row_num,
                    column=col_num,
                    value=value
                )

        for column in sheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column[0].column
            )

            for cell in column:

                if cell.value is not None:

                    length = len(
                        str(cell.value)
                    )

                    max_length = max(
                        max_length,
                        length
                    )

            sheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 3,
                40
            )

        sheet.freeze_panes = "A2"

    # =========================================================
    # COMPLETE EXCEL REPORT
    # =========================================================

    def export_excel(self):

        file_path = filedialog.asksaveasfilename(
            title="Save Complete Excel Report",
            defaultextension=".xlsx",
            filetypes=[
                ("Excel Files", "*.xlsx")
            ],
            initialfile="Student_Management_Complete_Report.xlsx"
        )

        if not file_path:
            return

        try:

            workbook = Workbook()

            workbook.remove(
                workbook.active
            )

            reports = [

                (
                    "Students",
                    """
                    SELECT
                        id,
                        student_id,
                        full_name,
                        email,
                        phone,
                        gender,
                        date_of_birth,
                        course,
                        semester,
                        address,
                        created_at
                    FROM students
                    """
                ),

                (
                    "Teachers",
                    """
                    SELECT
                        id,
                        teacher_id,
                        full_name,
                        email,
                        phone,
                        gender,
                        department,
                        designation,
                        subject,
                        qualification,
                        address,
                        created_at
                    FROM teachers
                    """
                ),

                (
                    "Courses",
                    """
                    SELECT
                        id,
                        course_code,
                        course_name,
                        duration,
                        description,
                        created_at
                    FROM courses
                    """
                ),

                (
                    "Subjects",
                    """
                    SELECT
                        id,
                        subject_code,
                        subject_name,
                        course_code,
                        semester,
                        teacher_id,
                        credits,
                        description,
                        created_at
                    FROM subjects
                    """
                ),

                (
                    "Enrollments",
                    """
                    SELECT
                        id,
                        student_id,
                        course_code,
                        enrollment_date,
                        semester,
                        status,
                        created_at
                    FROM enrollments
                    """
                ),

                (
                    "Attendance",
                    """
                    SELECT
                        id,
                        student_id,
                        subject_code,
                        attendance_date,
                        status,
                        created_at
                    FROM attendance
                    """
                ),

                (
                    "Marks",
                    """
                    SELECT
                        id,
                        student_id,
                        subject_code,
                        internal_marks,
                        external_marks,
                        total_marks,
                        grade,
                        created_at
                    FROM marks
                    """
                ),

                (
                    "Timetable",
                    """
                    SELECT
                        id,
                        course_code,
                        subject_code,
                        teacher_id,
                        day_of_week,
                        start_time,
                        end_time,
                        room_number,
                        semester,
                        created_at
                    FROM timetable
                    """
                ),

                (
                    "Announcements",
                    """
                    SELECT
                        id,
                        title,
                        message,
                        target_audience,
                        announcement_date,
                        created_at
                    FROM announcements
                    """
                )
            ]

            for sheet_name, query in reports:

                columns, rows = self.get_table_data(
                    query
                )

                if columns is not None:

                    self.add_sheet_to_workbook(
                        workbook,
                        sheet_name,
                        columns,
                        rows
                    )

            workbook.save(
                file_path
            )

            messagebox.showinfo(
                "Export Successful",
                "Complete Excel report downloaded successfully."
            )

        except Exception as e:

            messagebox.showerror(
                "Export Error",
                str(e)
            )

    # =========================================================
    # ATTENDANCE EXCEL
    # =========================================================

    def export_attendance_excel(self):

        file_path = filedialog.asksaveasfilename(
            title="Save Attendance Report",
            defaultextension=".xlsx",
            filetypes=[
                ("Excel Files", "*.xlsx")
            ],
            initialfile="Attendance_Report.xlsx"
        )

        if not file_path:
            return

        query = """
            SELECT
                a.id,
                a.student_id,
                s.full_name AS student_name,
                a.subject_code,
                a.attendance_date,
                a.status
            FROM attendance a
            LEFT JOIN students s
                ON a.student_id = s.student_id
            ORDER BY a.attendance_date DESC
        """

        columns, rows = self.get_table_data(
            query
        )

        if columns is None:
            return

        try:

            workbook = Workbook()

            self.add_sheet_to_workbook(
                workbook,
                "Attendance",
                columns,
                rows
            )

            workbook.save(
                file_path
            )

            messagebox.showinfo(
                "Export Successful",
                "Attendance report downloaded successfully."
            )

        except Exception as e:

            messagebox.showerror(
                "Export Error",
                str(e)
            )

    # =========================================================
    # MARKS EXCEL
    # =========================================================

    def export_marks_excel(self):

        file_path = filedialog.asksaveasfilename(
            title="Save Marks Report",
            defaultextension=".xlsx",
            filetypes=[
                ("Excel Files", "*.xlsx")
            ],
            initialfile="Marks_Report.xlsx"
        )

        if not file_path:
            return

        query = """
            SELECT
                m.id,
                m.student_id,
                s.full_name AS student_name,
                m.subject_code,
                m.internal_marks,
                m.external_marks,
                m.total_marks,
                m.grade
            FROM marks m
            LEFT JOIN students s
                ON m.student_id = s.student_id
            ORDER BY m.student_id
        """

        columns, rows = self.get_table_data(
            query
        )

        if columns is None:
            return

        try:

            workbook = Workbook()

            self.add_sheet_to_workbook(
                workbook,
                "Marks",
                columns,
                rows
            )

            workbook.save(
                file_path
            )

            messagebox.showinfo(
                "Export Successful",
                "Marks report downloaded successfully."
            )

        except Exception as e:

            messagebox.showerror(
                "Export Error",
                str(e)
            )

    # =========================================================
    # COMPLETE PDF REPORT
    # =========================================================

    def export_pdf(self):

        file_path = filedialog.asksaveasfilename(
            title="Save Complete PDF Report",
            defaultextension=".pdf",
            filetypes=[
                ("PDF Files", "*.pdf")
            ],
            initialfile="Student_Management_Complete_Report.pdf"
        )

        if not file_path:
            return

        connection = create_connection()

        if connection is None:

            messagebox.showerror(
                "Database Error",
                "Unable to connect to MySQL."
            )

            return

        try:

            cursor = connection.cursor()

            # COUNTS

            cursor.execute(
                "SELECT COUNT(*) FROM students"
            )
            total_students = cursor.fetchone()[0]

            cursor.execute(
                "SELECT COUNT(*) FROM teachers"
            )
            total_teachers = cursor.fetchone()[0]

            cursor.execute(
                "SELECT COUNT(*) FROM courses"
            )
            total_courses = cursor.fetchone()[0]

            cursor.execute(
                "SELECT COUNT(*) FROM subjects"
            )
            total_subjects = cursor.fetchone()[0]

            # ATTENDANCE

            cursor.execute(
                """
                SELECT
                    SUM(
                        CASE
                            WHEN status = 'Present'
                            THEN 1
                            ELSE 0
                        END
                    ),
                    COUNT(*)
                FROM attendance
                """
            )

            attendance_result = cursor.fetchone()

            present = attendance_result[0] or 0
            total_attendance = attendance_result[1] or 0

            if total_attendance > 0:

                attendance_percentage = (
                    present /
                    total_attendance
                ) * 100

            else:

                attendance_percentage = 0

            # PDF

            document = SimpleDocTemplate(
                file_path,
                pagesize=landscape(A4),
                rightMargin=30,
                leftMargin=30,
                topMargin=30,
                bottomMargin=30
            )

            styles = getSampleStyleSheet()

            title_style = ParagraphStyle(
                "ReportTitle",
                parent=styles["Title"],
                alignment=TA_CENTER,
                fontSize=22,
                spaceAfter=10
            )

            subtitle_style = ParagraphStyle(
                "Subtitle",
                parent=styles["Normal"],
                alignment=TA_CENTER,
                fontSize=10,
                spaceAfter=20
            )

            heading_style = ParagraphStyle(
                "Heading",
                parent=styles["Heading2"],
                fontSize=15,
                spaceBefore=10,
                spaceAfter=8
            )

            story = []

            story.append(
                Paragraph(
                    "STUDENT MANAGEMENT SYSTEM",
                    title_style
                )
            )

            story.append(
                Paragraph(
                    "Complete Reports & Analytics",
                    subtitle_style
                )
            )

            story.append(
                Paragraph(
                    "Generated on: "
                    + datetime.now().strftime(
                        "%d-%m-%Y %H:%M"
                    ),
                    subtitle_style
                )
            )

            # SUMMARY

            story.append(
                Paragraph(
                    "System Summary",
                    heading_style
                )
            )

            summary_data = [
                [
                    "Total Students",
                    "Total Teachers",
                    "Total Courses",
                    "Total Subjects",
                    "Attendance"
                ],
                [
                    str(total_students),
                    str(total_teachers),
                    str(total_courses),
                    str(total_subjects),
                    f"{attendance_percentage:.1f}%"
                ]
            ]

            summary_table = Table(
                summary_data,
                colWidths=[
                    130,
                    130,
                    130,
                    130,
                    130
                ]
            )

            summary_table.setStyle(
                TableStyle([
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.lightgrey
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold"
                    ),
                    (
                        "ALIGN",
                        (0, 0),
                        (-1, -1),
                        "CENTER"
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        10
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        8
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        8
                    )
                ])
            )

            story.append(
                summary_table
            )

            story.append(
                Spacer(1, 20)
            )

            # STUDENT REPORT

            story.append(
                Paragraph(
                    "Student Report",
                    heading_style
                )
            )

            cursor.execute(
                """
                SELECT
                    student_id,
                    full_name,
                    email,
                    phone,
                    course,
                    semester
                FROM students
                ORDER BY student_id
                """
            )

            students = cursor.fetchall()

            student_data = [
                [
                    "Student ID",
                    "Name",
                    "Email",
                    "Phone",
                    "Course",
                    "Semester"
                ]
            ]

            for row in students:

                student_data.append([
                    str(row[0] or ""),
                    str(row[1] or ""),
                    str(row[2] or ""),
                    str(row[3] or ""),
                    str(row[4] or ""),
                    str(row[5] or "")
                ])

            if len(student_data) > 1:

                student_table = Table(
                    student_data,
                    repeatRows=1
                )

                student_table.setStyle(
                    TableStyle([
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.lightgrey
                        ),
                        (
                            "FONTNAME",
                            (0, 0),
                            (-1, 0),
                            "Helvetica-Bold"
                        ),
                        (
                            "GRID",
                            (0, 0),
                            (-1, -1),
                            0.5,
                            colors.grey
                        ),
                        (
                            "FONTSIZE",
                            (0, 0),
                            (-1, -1),
                            8
                        )
                    ])
                )

                story.append(
                    student_table
                )

            else:

                story.append(
                    Paragraph(
                        "No student records available.",
                        styles["Normal"]
                    )
                )

            story.append(
                PageBreak()
            )

            # ATTENDANCE REPORT

            story.append(
                Paragraph(
                    "Attendance Report",
                    heading_style
                )
            )

            cursor.execute(
                """
                SELECT
                    a.student_id,
                    s.full_name,
                    a.subject_code,
                    a.attendance_date,
                    a.status
                FROM attendance a
                LEFT JOIN students s
                    ON a.student_id = s.student_id
                ORDER BY a.attendance_date DESC
                """
            )

            attendance_rows = cursor.fetchall()

            attendance_data = [
                [
                    "Student ID",
                    "Student Name",
                    "Subject",
                    "Date",
                    "Status"
                ]
            ]

            for row in attendance_rows:

                attendance_data.append([
                    str(row[0] or ""),
                    str(row[1] or ""),
                    str(row[2] or ""),
                    str(row[3] or ""),
                    str(row[4] or "")
                ])

            if len(attendance_data) > 1:

                attendance_table = Table(
                    attendance_data,
                    repeatRows=1
                )

                attendance_table.setStyle(
                    TableStyle([
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.lightgrey
                        ),
                        (
                            "FONTNAME",
                            (0, 0),
                            (-1, 0),
                            "Helvetica-Bold"
                        ),
                        (
                            "GRID",
                            (0, 0),
                            (-1, -1),
                            0.5,
                            colors.grey
                        ),
                        (
                            "ALIGN",
                            (0, 0),
                            (-1, -1),
                            "CENTER"
                        ),
                        (
                            "FONTSIZE",
                            (0, 0),
                            (-1, -1),
                            8
                        )
                    ])
                )

                story.append(
                    attendance_table
                )

            else:

                story.append(
                    Paragraph(
                        "No attendance records available.",
                        styles["Normal"]
                    )
                )

            story.append(
                PageBreak()
            )

            # MARKS REPORT

            story.append(
                Paragraph(
                    "Marks & Grade Report",
                    heading_style
                )
            )

            cursor.execute(
                """
                SELECT
                    m.student_id,
                    s.full_name,
                    m.subject_code,
                    m.internal_marks,
                    m.external_marks,
                    m.total_marks,
                    m.grade
                FROM marks m
                LEFT JOIN students s
                    ON m.student_id = s.student_id
                ORDER BY m.student_id
                """
            )

            marks_rows = cursor.fetchall()

            marks_data = [
                [
                    "Student ID",
                    "Student Name",
                    "Subject",
                    "Internal",
                    "External",
                    "Total",
                    "Grade"
                ]
            ]

            for row in marks_rows:

                marks_data.append([
                    str(row[0] or ""),
                    str(row[1] or ""),
                    str(row[2] or ""),
                    str(row[3] or ""),
                    str(row[4] or ""),
                    str(row[5] or ""),
                    str(row[6] or "")
                ])

            if len(marks_data) > 1:

                marks_table = Table(
                    marks_data,
                    repeatRows=1
                )

                marks_table.setStyle(
                    TableStyle([
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.lightgrey
                        ),
                        (
                            "FONTNAME",
                            (0, 0),
                            (-1, 0),
                            "Helvetica-Bold"
                        ),
                        (
                            "GRID",
                            (0, 0),
                            (-1, -1),
                            0.5,
                            colors.grey
                        ),
                        (
                            "ALIGN",
                            (0, 0),
                            (-1, -1),
                            "CENTER"
                        ),
                        (
                            "FONTSIZE",
                            (0, 0),
                            (-1, -1),
                            8
                        )
                    ])
                )

                story.append(
                    marks_table
                )

            else:

                story.append(
                    Paragraph(
                        "No marks records available.",
                        styles["Normal"]
                    )
                )

            document.build(
                story
            )

            cursor.close()
            connection.close()

            messagebox.showinfo(
                "Export Successful",
                "Complete PDF report downloaded successfully."
            )

        except Exception as e:

            connection.close()

            messagebox.showerror(
                "PDF Export Error",
                str(e)
            )

    # =========================================================
    # REFRESH
    # =========================================================

    def refresh_reports(self):

        self.load_statistics()


# =============================================================
# STANDALONE TEST
# =============================================================

if __name__ == "__main__":

    ctk.set_appearance_mode("light")
    ctk.set_default_color_theme("blue")

    app = ctk.CTk()

    app.title(
        "Reports & Analytics"
    )

    app.geometry(
        "1450x800"
    )

    page = ReportsPage(app)

    page.pack(
        fill="both",
        expand=True
    )

    app.mainloop()